from io import BytesIO
from datetime import date, datetime
from typing import Any


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return "; ".join(f"{key}: {item}" for key, item in value.items()) if isinstance(value, dict) else ", ".join(map(str, value))
    return str(value)


def build_excel(report: dict) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append([report["title"]])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([f"From: {_display(report.get('date_from'))}", f"To: {_display(report.get('date_to'))}"])
    sheet.append([])
    sheet.append(report["columns"])
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
    for row in report["rows"]:
        sheet.append([_display(row.get(column)) for column in report["columns"]])
    for column in sheet.columns:
        width = min(max(max(len(_display(cell.value)) for cell in column) + 2, 12), 36)
        sheet.column_dimensions[column[0].column_letter].width = width

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary["A1"].font = summary["B1"].font = Font(bold=True)
    for key, value in report.get("summary", {}).items():
        summary.append([key.replace("_", " ").title(), _display(value)])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_pdf(report: dict) -> BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    story = [Paragraph(report["title"], styles["Title"]), Paragraph(f"From: {_display(report.get('date_from'))} &nbsp;&nbsp; To: {_display(report.get('date_to'))}", styles["Normal"]), Spacer(1, 12)]
    data = [[Paragraph(str(column).replace("_", " ").title(), styles["Normal"]) for column in report["columns"]]]
    data.extend([[_display(row.get(column)) for column in report["columns"]] for row in report["rows"]])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    story.append(table)
    story.append(Spacer(1, 12))
    summary_lines = [f"{key.replace('_', ' ').title()}: {_display(value)}" for key, value in report.get("summary", {}).items()]
    story.append(Paragraph("Summary", styles["Heading2"]))
    story.extend(Paragraph(line, styles["Normal"]) for line in summary_lines)
    document.build(story)
    output.seek(0)
    return output
